# -*- coding: utf-8 -*-
"""표준 발주 양식(최종 통합본) 정의 및 스타일링.

최종양식(12열): NO · 주문인명 · 주문인핸드폰번호 · 수령인명 · 수령인핸드폰번호 ·
우편번호 · 주소 · 배송메세지 · 상품정보 · 주문수량 · 택배사 · 송장번호
- 글꼴: 맑은 고딕 10pt, 전 셀 얇은 테두리, 헤더 연한 회색 채움
"""
from __future__ import annotations

from copy import copy

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ---------------------------------------------------------------------------
# 표준 양식 열 정의 (기준: 최종양식 12열)
# ---------------------------------------------------------------------------
STANDARD_COLUMNS = [
    ("no",             "NO"),
    ("orderer",        "주문인명"),
    ("orderer_phone",  "주문인핸드폰번호"),
    ("recipient",      "수령인명"),
    ("recipient_phone", "수령인핸드폰번호"),
    ("postcode",       "우편번호"),
    ("address",        "주소"),
    ("message",        "배송메세지"),
    ("product",        "상품정보"),
    ("qty",            "주문수량"),
    ("courier",        "택배사"),
    ("invoice",        "송장번호"),
]

FIELD_KEYS = [key for key, _ in STANDARD_COLUMNS]

# 열 너비
COLUMN_WIDTHS = {
    "A": 4.25, "B": 10.0, "C": 15.0, "D": 10.0, "E": 15.0, "F": 8.0,
    "G": 32.0, "H": 16.0, "I": 40.0, "J": 8.0, "K": 10.0, "L": 15.0,
}

# 가운데 정렬 열(1-based): NO, 우편번호, 주문수량
_CENTER_COLS = {1, 6, 10}
# 텍스트 서식 열(앞자리 0/서식 보존): 주문인핸드폰, 수령인핸드폰, 우편번호, 송장번호
_TEXT_COLS = {3, 5, 6, 12}

FONT_NAME = "맑은 고딕"
FONT_SIZE = 10

_THIN = Side(style="thin", color="FFBFBFBF")
THIN_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
HEADER_FILL = PatternFill(patternType="solid", fgColor="FFF2F2F2")


def write_standard_sheet(ws, rows):
    """``rows`` (표준 필드 dict 리스트)를 표준 서식으로 ``ws`` 에 기록한다."""
    base_font = Font(name=FONT_NAME, size=FONT_SIZE)

    # 헤더 ------------------------------------------------------------------
    for c, (_key, header) in enumerate(STANDARD_COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font = copy(base_font)
        cell.border = THIN_BORDER
        cell.fill = copy(HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 데이터 ----------------------------------------------------------------
    for r, row in enumerate(rows, start=2):
        for c, key in enumerate(FIELD_KEYS, start=1):
            cell = ws.cell(row=r, column=c, value=row.get(key))
            cell.font = copy(base_font)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                horizontal="center" if c in _CENTER_COLS else "left",
                vertical="center",
            )
            if c in _TEXT_COLS:
                cell.number_format = "@"

    # 열 너비 ---------------------------------------------------------------
    for col_letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    return ws
