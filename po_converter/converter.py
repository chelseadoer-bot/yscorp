# -*- coding: utf-8 -*-
"""발주 파일 통합 변환기.

업체 원본 파일 1개를 받아 표준 발주 양식(최종 통합본)으로 변환한다.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

import openpyxl

from .cleaning import clean_text
from .format_spec import write_standard_sheet
from .vendors import Row, Vendor, detect_vendor, normalize_header


@dataclass
class ConversionResult:
    vendor: Vendor
    vendor_name: str          # 파일명에 사용할 업체명
    name_confident: bool      # 업체명을 원본에서 확정했는지
    date: Optional[str]       # YYMMDD (원본에서 추출한 경우)
    rows: List[dict]          # 표준 필드 dict 리스트
    source_sheet: str


def _find_header_row(ws, max_scan: int = 5):
    """상단 몇 행을 훑어 업체 시그니처와 일치하는 헤더 행을 찾는다.

    반환: (vendor, header_row_index, header_map) 또는 (None, None, None).
    """
    for r in range(1, min(ws.max_row, max_scan) + 1):
        header_map = {}
        present = set()
        for c in range(1, ws.max_column + 1):
            norm = normalize_header(ws.cell(r, c).value)
            if norm:
                header_map.setdefault(norm, c)
                present.add(norm)
        vendor = detect_vendor(present)
        if vendor:
            return vendor, r, header_map
    return None, None, None


def _to_yymmdd(value) -> Optional[str]:
    text = clean_text(value)
    m = re.search(r"(20)?(\d{2})(\d{2})(\d{2})", text.replace("-", "").replace(".", ""))
    if m:
        return f"{m.group(2)}{m.group(3)}{m.group(4)}"
    return None


def _date_from_filename(path: str) -> Optional[str]:
    """파일명에 들어있는 날짜를 YYMMDD로 추출.

    - YYYYMMDD / YYMMDD (예: 20260730, 260730) → 그대로
    - MMDD (예: 0730) → 올해 연도(YY)를 앞에 붙임 → 260730
    """
    import datetime

    name = Path(path).stem
    m = re.search(r"(?:20)?(\d{2})(\d{2})(\d{2})", name)
    if m:
        return f"{m.group(1)}{m.group(2)}{m.group(3)}"
    m = re.search(r"(?<!\d)(\d{2})(\d{2})(?!\d)", name)
    if m and 1 <= int(m.group(1)) <= 12 and 1 <= int(m.group(2)) <= 31:
        return f"{datetime.date.today():%y}{m.group(1)}{m.group(2)}"
    return None


def convert_workbook(path: str, sheet: Optional[str] = None) -> ConversionResult:
    """원본 파일을 읽어 표준 행 리스트로 변환한다."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.active

    # 데이터가 있는 첫 시트를 자동 선택(빈 시트 건너뛰기)
    if ws.max_row <= 1 and not sheet:
        for cand in wb.worksheets:
            if cand.max_row > 1:
                ws = cand
                break

    vendor, header_row, header_map = _find_header_row(ws)
    if vendor is None:
        raise ValueError(
            f"업체 양식을 인식하지 못했습니다: {Path(path).name}\n"
            f"  인식된 헤더: {[ws.cell(1, c).value for c in range(1, ws.max_column + 1)]}"
        )

    # 데이터 행 → 표준 행
    rows: List[dict] = []
    order_date: Optional[str] = None
    for r in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if not any(clean_text(v) for v in values):
            continue  # 완전히 빈 행 건너뛰기
        row = Row(values, header_map)
        std = vendor.map_row(row)
        # 이름/주소가 모두 비어 있으면 유효 주문이 아님
        if not clean_text(std.get("recipient")) and not clean_text(std.get("address")):
            continue
        std["no"] = len(rows) + 1
        rows.append(std)

        if order_date is None and vendor.date_header:
            order_date = _to_yymmdd(row.h(vendor.date_header))

    # 날짜: 데이터 > 파일명 순으로 확정 (없으면 호출부에서 오늘 날짜 사용)
    if order_date is None:
        order_date = _date_from_filename(path)

    # 업체명 확정
    vendor_name = vendor.default_name
    if vendor.name_from_headers:
        vendor_name = vendor.name_from_headers(header_map) or vendor_name
    if vendor.name_from_sheet:
        vendor_name = vendor.name_from_sheet(ws.title) or vendor_name
    if vendor.name_from_filename:
        vendor_name = vendor.name_from_filename(Path(path).stem) or vendor_name

    return ConversionResult(
        vendor=vendor,
        vendor_name=vendor_name,
        name_confident=vendor.name_confident,
        date=order_date,
        rows=rows,
        source_sheet=ws.title,
    )


def write_output(rows: List[dict], out_path: str, sheet_title: str = "발주") -> str:
    """표준 행 리스트를 표준 양식 xlsx 파일로 저장한다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] if sheet_title else "발주"
    write_standard_sheet(ws, rows)
    wb.save(out_path)
    return out_path


def build_output_filename(date: str, vendor_name: str) -> str:
    """파일명 규칙: '날짜_업체명_발주 수정본.xlsx'."""
    return f"{date}_{vendor_name}_발주 수정본.xlsx"
